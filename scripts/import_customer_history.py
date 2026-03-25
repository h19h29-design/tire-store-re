from __future__ import annotations

import argparse
import re
import sqlite3
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


SERVICE_KEYWORDS = (
    "alignment",
    "얼라이",
    "장착",
    "오일",
    "필터",
    "교환",
    "정비",
    "보관",
    "수리",
    "중고",
)


@dataclass
class ParsedHistoryRow:
    sold_at: str
    customer_name: str
    phone: str
    normalized_phone: str
    vehicle_model: str
    plate_number: str
    normalized_plate_number: str
    odometer: int
    pattern: str
    size_label: str
    quantity: int
    alignment_amount: int
    total_amount: int
    memo: str
    line_type: str


def normalize_text(value: object) -> str:
    text = str(value or "").strip().lower()
    for token in (" ", "-", "/", ".", "(", ")", "\n", "\r", "\t"):
        text = text.replace(token, "")
    return text


def sanitize_text(value: object) -> str:
    text = str(value or "").strip()
    return "" if text in {"", "0", "-", "None"} else text


def sanitize_phone(value: object) -> str:
    raw = sanitize_text(value)
    digits = re.sub(r"\D", "", raw)
    return raw if len(digits) >= 8 else ""


def normalize_phone(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def sanitize_plate(value: object) -> str:
    raw = sanitize_text(value)
    normalized = normalize_plate(raw)
    digit_count = sum(ch.isdigit() for ch in normalized)
    return raw if digit_count >= 4 and len(normalized) >= 6 else ""


def normalize_plate(value: object) -> str:
    return re.sub(r"[\s-]+", "", str(value or "")).lower()


def parse_number(value: object) -> int:
    text = str(value or "").strip().replace(",", "").replace(" ", "")
    if text in {"", "-", "None"}:
        return 0
    try:
        return int(float(text))
    except ValueError:
        digits = re.sub(r"\D", "", text)
        if not digits:
            return 0
        return int(digits)


def parse_thousand_won(value: object) -> int:
    text = str(value or "").strip().replace(",", "")
    if text in {"", "-", "None"}:
        return 0
    try:
        return round(float(text) * 1000)
    except ValueError:
        return 0


def parse_note_amount_won(value: str) -> int:
    match = re.search(r"(\d[\d,]*)\s*원", value)
    if not match:
        return 0
    return int(match.group(1).replace(",", ""))


def parse_odometer(value: object) -> int:
    number = parse_number(value)
    return number if number >= 1000 else 0


def combine_text_cells(values: Iterable[object]) -> str:
    return " / ".join(filter(None, (sanitize_text(value) for value in values)))


def looks_like_tire_size(value: str) -> bool:
    return sum(len(part) for part in re.findall(r"\d+", value)) >= 5 and len(re.findall(r"\d+", value)) >= 2


def looks_like_tire_pattern(value: str) -> bool:
    text = sanitize_text(value)
    if not text or text.isdigit() or is_service_related(text):
        return False
    return bool(re.search(r"[A-Za-z가-힣]", text))


def is_service_related(*values: str) -> bool:
    combined = normalize_text(" ".join(values))
    return any(normalize_text(keyword) in combined for keyword in SERVICE_KEYWORDS)


def extract_date(raw: object, current: str) -> str:
    digits = "".join(ch for ch in str(raw or "") if ch.isdigit())
    if len(digits) < 8:
        return current
    year = int(digits[0:4])
    month = int(digits[4:6]) if 1 <= int(digits[4:6]) <= 12 else 1
    day = int(digits[6:8]) if 1 <= int(digits[6:8]) <= 31 else 1
    return f"{year:04d}-{month:02d}-{day:02d}"


def detect_layout(values: list[object]) -> str | None:
    normalized = [normalize_text(value) for value in values]
    has_plate = "차량번호" in normalized
    has_phone = "연락처" in normalized
    if not has_plate or not has_phone:
        return None
    if "날자" in normalized or "날짜" in normalized or "이름" in normalized:
        return "current"
    return "legacy"


def classify_line_type(pattern: str, size_label: str, quantity: int, alignment_amount: int, memo: str) -> str:
    if quantity > 0 and (looks_like_tire_size(size_label) or looks_like_tire_pattern(pattern)) and not is_service_related(pattern, size_label):
        return "tire"
    if alignment_amount > 0 or is_service_related(pattern, size_label, memo):
        return "service"
    if quantity <= 0 and (pattern.strip() or size_label.strip() or memo.strip()):
        return "service"
    if quantity > 0 and (looks_like_tire_size(size_label) or looks_like_tire_pattern(pattern)):
        return "tire"
    return "other"


def build_service_name(pattern: str, size_label: str, memo: str) -> str:
    for candidate in (pattern, size_label, memo):
        text = sanitize_text(candidate)
        if text and not looks_like_tire_size(text):
            return text
    return "추가 작업"


def parse_customer_history_rows(workbook_path: Path) -> list[ParsedHistoryRow]:
    workbook = load_workbook(workbook_path, data_only=True)
    customer_sheet_name = next((name for name in workbook.sheetnames if "고객등록" in name), workbook.sheetnames[4])
    sheet = workbook[customer_sheet_name]

    current_layout = "legacy"
    current_date = ""
    parsed_rows: list[ParsedHistoryRow] = []

    for row in sheet.iter_rows(values_only=True):
        values = list(row[:13]) + [None] * max(0, 13 - len(row))
        next_layout = detect_layout(values)
        if next_layout:
            current_layout = next_layout
            continue

        current_date = extract_date(values[0], current_date)
        if not current_date:
            continue

        if current_layout == "legacy":
            customer_name = ""
            odometer = max(parse_odometer(values[1]), parse_odometer(values[0]))
            phone = sanitize_phone(values[2]) or sanitize_phone(values[1])
            vehicle_model = sanitize_text(values[3])
            plate_number = sanitize_plate(values[4])
            raw_pattern = sanitize_text(values[5])
            raw_size = sanitize_text(values[6])
            if looks_like_tire_size(raw_pattern) and not looks_like_tire_size(raw_size):
                pattern = ""
                size_label = raw_pattern
                quantity = parse_number(values[6])
                alignment_amount = parse_thousand_won(values[7])
            else:
                pattern = raw_pattern
                size_label = raw_size
                quantity = parse_number(values[7])
                alignment_amount = parse_thousand_won(values[8])
            total_amount = parse_thousand_won(values[10]) or parse_thousand_won(values[9])
            memo = combine_text_cells((values[11], values[12]))
        else:
            odometer = parse_odometer(values[1])
            phone = sanitize_phone(values[2]) or sanitize_phone(values[1])
            customer_name = ""
            if phone:
                if odometer == 0:
                    customer_name = sanitize_text(values[1])
            else:
                customer_name = sanitize_text(values[2]) or sanitize_text(values[1])
            vehicle_model = sanitize_text(values[3])
            plate_number = sanitize_plate(values[4])
            raw_pattern = sanitize_text(values[5])
            raw_size = sanitize_text(values[6])
            if looks_like_tire_size(raw_pattern) and not looks_like_tire_size(raw_size):
                pattern = ""
                size_label = raw_pattern
                quantity = parse_number(values[6])
                alignment_amount = parse_thousand_won(values[7])
            elif raw_pattern and not looks_like_tire_pattern(raw_pattern) and parse_number(values[6]) > 0:
                pattern = raw_pattern
                size_label = ""
                quantity = parse_number(values[6])
                alignment_amount = parse_thousand_won(values[7])
            else:
                pattern = raw_pattern
                size_label = raw_size
                quantity = parse_number(values[7])
                alignment_amount = parse_thousand_won(values[8])
            total_amount = parse_thousand_won(values[10])
            memo = combine_text_cells((values[11], values[12]))

        if total_amount <= 0:
            total_amount = parse_note_amount_won(memo)

        line_type = classify_line_type(pattern, size_label, quantity, alignment_amount, memo)
        normalized_phone = normalize_phone(phone)
        normalized_plate = normalize_plate(plate_number)

        if not any(
            [
                customer_name,
                phone,
                vehicle_model,
                plate_number,
                pattern,
                size_label,
                quantity,
                alignment_amount,
                total_amount,
                memo,
            ]
        ):
            continue

        if quantity <= 0 and alignment_amount <= 0 and total_amount <= 0:
            continue

        if line_type == "other" and total_amount <= 0 and alignment_amount <= 0:
            continue

        parsed_rows.append(
            ParsedHistoryRow(
                sold_at=current_date,
                customer_name=customer_name,
                phone=phone,
                normalized_phone=normalized_phone,
                vehicle_model=vehicle_model,
                plate_number=plate_number,
                normalized_plate_number=normalized_plate,
                odometer=odometer,
                pattern=pattern,
                size_label=size_label,
                quantity=max(quantity, 0),
                alignment_amount=max(alignment_amount, 0),
                total_amount=max(total_amount, 0),
                memo=memo,
                line_type=line_type,
            )
        )

    return parsed_rows


def load_existing_sale_keys(conn: sqlite3.Connection) -> set[str]:
    rows = conn.execute(
        """
        SELECT
          SUBSTR(sales.sold_at, 1, 10) AS sold_at,
          COALESCE(vehicles.normalized_plate_number, '') AS normalized_plate_number,
          COALESCE(sale_lines.item_snapshot_name, '') AS item_snapshot_name,
          COALESCE(sale_lines.size_snapshot, '') AS size_snapshot,
          COALESCE(sale_lines.quantity, 0) AS quantity,
          COALESCE(sale_lines.line_total, 0) AS line_total,
          COALESCE(sale_lines.memo, '') AS memo,
          COALESCE(sale_lines.line_type, '') AS line_type
        FROM sales
        INNER JOIN sale_lines
          ON sale_lines.sale_id = sales.id
        LEFT JOIN vehicles
          ON vehicles.id = sales.vehicle_id
        """
    ).fetchall()
    return {
        build_sale_key(
            sold_at=row[0],
            normalized_plate_number=row[1],
            snapshot_name=row[2],
            size_snapshot=row[3],
            quantity=int(row[4] or 0),
            line_total=int(row[5] or 0),
            memo=row[6] or "",
            line_type=row[7] or "",
        )
        for row in rows
    }


def build_sale_key(
    *,
    sold_at: str,
    normalized_plate_number: str,
    snapshot_name: str,
    size_snapshot: str,
    quantity: int,
    line_total: int,
    memo: str,
    line_type: str,
) -> str:
    return "|".join(
        [
            sold_at,
            normalized_plate_number,
            normalize_text(snapshot_name),
            normalize_text(size_snapshot),
            str(quantity),
            str(line_total),
            normalize_text(memo),
            normalize_text(line_type),
        ]
    )


def ensure_customer(conn: sqlite3.Connection, row: ParsedHistoryRow) -> int | None:
    if row.normalized_phone:
        existing = conn.execute(
            "SELECT id FROM customers WHERE normalized_phone = ? LIMIT 1",
            (row.normalized_phone,),
        ).fetchone()
        if existing:
            conn.execute(
                """
                UPDATE customers
                SET
                  name = CASE WHEN TRIM(name) = '' AND ? <> '' THEN ? ELSE name END,
                  phone = CASE WHEN TRIM(phone) = '' AND ? <> '' THEN ? ELSE phone END,
                  memo = CASE WHEN TRIM(COALESCE(memo, '')) = '' AND ? <> '' THEN ? ELSE memo END,
                  updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (row.customer_name, row.customer_name, row.phone, row.phone, row.memo, row.memo, existing[0]),
            )
            return int(existing[0])

    if not row.customer_name and not row.phone and not row.normalized_plate_number:
        return None

    inserted = conn.execute(
        """
        INSERT INTO customers (name, phone, normalized_phone, memo, updated_at)
        VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (row.customer_name, row.phone, row.normalized_phone, row.memo),
    )
    return int(inserted.lastrowid)


def ensure_vehicle(conn: sqlite3.Connection, row: ParsedHistoryRow, customer_id: int | None) -> int | None:
    if row.normalized_plate_number:
        existing = conn.execute(
            "SELECT id, odometer FROM vehicles WHERE normalized_plate_number = ? LIMIT 1",
            (row.normalized_plate_number,),
        ).fetchone()
        if existing:
            conn.execute(
                """
                UPDATE vehicles
                SET
                  customer_id = COALESCE(customer_id, ?),
                  plate_number = CASE WHEN TRIM(plate_number) = '' AND ? <> '' THEN ? ELSE plate_number END,
                  model_name = CASE WHEN TRIM(model_name) = '' AND ? <> '' THEN ? ELSE model_name END,
                  odometer = CASE WHEN ? > COALESCE(odometer, 0) THEN ? ELSE odometer END,
                  memo = CASE WHEN TRIM(COALESCE(memo, '')) = '' AND ? <> '' THEN ? ELSE memo END,
                  updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (
                    customer_id,
                    row.plate_number,
                    row.plate_number,
                    row.vehicle_model,
                    row.vehicle_model,
                    row.odometer,
                    row.odometer,
                    row.memo,
                    row.memo,
                    existing[0],
                ),
            )
            return int(existing[0])

    if not row.vehicle_model and not row.plate_number:
        return None

    inserted = conn.execute(
        """
        INSERT INTO vehicles (
          customer_id,
          plate_number,
          normalized_plate_number,
          model_name,
          odometer,
          memo,
          updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (
            customer_id,
            row.plate_number,
            row.normalized_plate_number,
            row.vehicle_model,
            row.odometer,
            row.memo,
        ),
    )
    return int(inserted.lastrowid)


def import_customer_history(workbook_path: Path, db_path: Path) -> tuple[int, int, int]:
    rows = parse_customer_history_rows(workbook_path)
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    existing_sale_keys = load_existing_sale_keys(conn)

    inserted_sales = 0
    inserted_lines = 0
    inserted_logs = 0

    try:
        for sequence, row in enumerate(rows, start=1):
            primary_total = max(row.total_amount - min(row.alignment_amount, row.total_amount), 0)
            sale_total = row.total_amount or row.alignment_amount
            primary_snapshot_name = row.pattern or (row.size_label if row.line_type != "tire" else "타이어")
            if row.line_type == "service":
                primary_snapshot_name = build_service_name(row.pattern, row.size_label, row.memo)

            sale_key = build_sale_key(
                sold_at=row.sold_at,
                normalized_plate_number=row.normalized_plate_number,
                snapshot_name=primary_snapshot_name,
                size_snapshot=row.size_label if row.line_type == "tire" else "",
                quantity=max(row.quantity, 1 if row.line_type == "service" and primary_total > 0 else 0),
                line_total=primary_total if primary_total > 0 else sale_total,
                memo=row.memo,
                line_type=row.line_type,
            )
            if sale_key in existing_sale_keys:
                continue

            customer_id = ensure_customer(conn, row)
            vehicle_id = ensure_vehicle(conn, row, customer_id)
            sale_number = f"HIS-{row.sold_at.replace('-', '')}-{sequence:05d}"
            sale_insert = conn.execute(
                """
                INSERT INTO sales (
                  sale_number,
                  sold_at,
                  customer_id,
                  vehicle_id,
                  total_amount,
                  card_amount,
                  naver_amount,
                  cash_amount,
                  card_fee_rate_basis_points,
                  is_naver_card,
                  memo,
                  updated_at
                ) VALUES (?, ?, ?, ?, ?, 0, 0, ?, 0, 0, ?, CURRENT_TIMESTAMP)
                """,
                (
                    sale_number,
                    f"{row.sold_at} 00:00:00",
                    customer_id,
                    vehicle_id,
                    sale_total,
                    sale_total,
                    row.memo,
                ),
            )
            sale_id = int(sale_insert.lastrowid)
            inserted_sales += 1

            if primary_total > 0 or row.line_type == "tire":
                primary_quantity = max(row.quantity, 1 if row.line_type == "service" and primary_total > 0 else 1)
                primary_line_total = primary_total if primary_total > 0 else sale_total
                primary_unit_price = round(primary_line_total / max(primary_quantity, 1))
                conn.execute(
                    """
                    INSERT INTO sale_lines (
                      sale_id,
                      line_type,
                      item_id,
                      item_snapshot_name,
                      size_snapshot,
                      cost_price_snapshot,
                      quantity,
                      unit_price,
                      line_total,
                      memo
                    ) VALUES (?, ?, NULL, ?, ?, NULL, ?, ?, ?, ?)
                    """,
                    (
                        sale_id,
                        row.line_type,
                        primary_snapshot_name or ("타이어" if row.line_type == "tire" else "추가 작업"),
                        row.size_label if row.line_type == "tire" else "",
                        primary_quantity,
                        primary_unit_price,
                        primary_line_total,
                        row.memo,
                    ),
                )
                inserted_lines += 1

                if row.line_type == "service":
                    conn.execute(
                        """
                        INSERT INTO work_logs (
                          sale_id,
                          customer_id,
                          vehicle_id,
                          work_type,
                          amount,
                          memo,
                          worked_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            sale_id,
                            customer_id,
                            vehicle_id,
                            primary_snapshot_name or "추가 작업",
                            primary_line_total,
                            row.memo,
                            f"{row.sold_at} 00:00:00",
                        ),
                    )
                    inserted_logs += 1

            if row.alignment_amount > 0:
                conn.execute(
                    """
                    INSERT INTO sale_lines (
                      sale_id,
                      line_type,
                      item_id,
                      item_snapshot_name,
                      size_snapshot,
                      cost_price_snapshot,
                      quantity,
                      unit_price,
                      line_total,
                      memo
                    ) VALUES (?, 'service', NULL, '얼라이먼트', '', NULL, 1, ?, ?, ?)
                    """,
                    (sale_id, row.alignment_amount, row.alignment_amount, row.memo),
                )
                conn.execute(
                    """
                    INSERT INTO work_logs (
                      sale_id,
                      customer_id,
                      vehicle_id,
                      work_type,
                      amount,
                      memo,
                      worked_at
                    ) VALUES (?, ?, ?, '얼라이먼트', ?, ?, ?)
                    """,
                    (
                        sale_id,
                        customer_id,
                        vehicle_id,
                        row.alignment_amount,
                        row.memo,
                        f"{row.sold_at} 00:00:00",
                    ),
                )
                inserted_lines += 1
                inserted_logs += 1

            existing_sale_keys.add(sale_key)

        conn.execute(
            "INSERT INTO imports (source_file, status, note) VALUES (?, 'success', ?)",
            (str(workbook_path), f"customer-history rows={len(rows)} inserted_sales={inserted_sales}"),
        )
        conn.commit()
    finally:
        conn.close()

    return inserted_sales, inserted_lines, inserted_logs


def main() -> None:
    parser = argparse.ArgumentParser(description="Import customer history from the inventory workbook.")
    parser.add_argument("workbook_path", help="Path to the inventory workbook (.xlsx)")
    parser.add_argument(
        "--db-path",
        default=str(Path.home() / "AppData" / "Roaming" / "com.tirestore.desktop" / "tire-store.db"),
        help="Path to tire-store.db",
    )
    args = parser.parse_args()

    workbook_path = Path(args.workbook_path).expanduser().resolve()
    db_path = Path(args.db_path).expanduser().resolve()

    inserted_sales, inserted_lines, inserted_logs = import_customer_history(workbook_path, db_path)
    print(
        f"customer_history_imported sales={inserted_sales} sale_lines={inserted_lines} work_logs={inserted_logs} db={db_path}"
    )


if __name__ == "__main__":
    main()
